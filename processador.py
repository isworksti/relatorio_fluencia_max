import os
import re
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

LEGENDAS_COMPLETAS = [
    ("P.L.N. 1", "Pré - Leitor Nível 1", "E5C99C", "000000"),
    ("P.L.N. 2", "Pré- Leitor Nível 2", "E3C700", "000000"),
    ("P.L.N. 3", "Pré- Leitor Nível 3", "F5D000", "000000"),
    ("P.L.N. 4", "Pré- Leitor Nível 4", "E5B800", "000000"),
    ("L.I.", "Leitor Iniciante", "9AF746", "000000"),
    ("L.F.", "Leitor Fluente", "00B626", "FFFFFF")
]

MAPA_DICT = {item[0]: item for item in LEGENDAS_COMPLETAS}
MAPA_DICT["P.L.N."] = ("P.L.N.", "Pré - Leitor Nível 1", "E5C99C", "000000")


def identificar_bimestre(conteudo_html):
    """Inspeciona o texto interno do HTML para identificar o bimestre."""
    padroes = [
        (r'1[º°o]?\s*Bimestre', "1º Bimestre"),
        (r'2[º°o]?\s*Bimestre', "2º Bimestre"),
        (r'3[º°o]?\s*Bimestre', "3º Bimestre"),
        (r'4[º°o]?\s*Bimestre', "4º Bimestre"),
        (r'Bimestre\s*:\s*1', "1º Bimestre"),
        (r'Bimestre\s*:\s*2', "2º Bimestre"),
        (r'Bimestre\s*:\s*3', "3º Bimestre"),
        (r'Bimestre\s*:\s*4', "4º Bimestre"),
    ]
    for padrao, bimestre in padroes:
        if re.search(padrao, conteudo_html, re.IGNORECASE):
            return bimestre
    return None


def processar_pasta_relatorios(pasta_origem):
    """Lê todos os arquivos HTML de uma pasta e retorna os dados estruturados."""
    dados_alunos = {}
    info_cabecalho = {"Escola": "", "Turma": "", "Série": "", "Ano Letivo": "", "Disciplina": ""}
    bimestres_encontrados = []

    arquivos_na_pasta = [f for f in os.listdir(pasta_origem) if f.lower().endswith(('.html', '.htm'))]

    for nome_arquivo in arquivos_na_pasta:
        caminho_completo = os.path.join(pasta_origem, nome_arquivo)
        
        conteudo_html = ""
        for enc in ["utf-8-sig", "utf-16", "utf-8", "latin-1"]:
            try:
                with open(caminho_completo, "r", encoding=enc) as f:
                    conteudo_html = f.read()
                if "<html" in conteudo_html.lower() or "<table" in conteudo_html.lower():
                    break
            except Exception:
                continue

        if not conteudo_html:
            continue

        bimestre = identificar_bimestre(conteudo_html)
        if not bimestre:
            continue
        
        if bimestre not in bimestres_encontrados:
            bimestres_encontrados.append(bimestre)

        soup = BeautifulSoup(conteudo_html, "html.parser")

        # Extração de dados de cabeçalho
        if not info_cabecalho["Escola"]:
            div_grid = soup.find("div", class_=lambda c: c and "grid-cols-3" in c)
            if div_grid:
                flex_items = div_grid.find_all("div", class_="flex")
                for item in flex_items:
                    texto_item = item.text.strip()
                    if "Escola:" in texto_item:
                        info_cabecalho["Escola"] = texto_item.replace("Escola:", "").strip()
                    elif "Turma:" in texto_item:
                        info_cabecalho["Turma"] = texto_item.replace("Turma:", "").strip()
                    elif "Série:" in texto_item:
                        info_cabecalho["Série"] = texto_item.replace("Série:", "").strip()
                    elif "Ano Letivo:" in texto_item:
                        info_cabecalho["Ano Letivo"] = texto_item.replace("Ano Letivo:", "").strip()
                    elif "Disciplina:" in texto_item:
                        info_cabecalho["Disciplina"] = texto_item.replace("Disciplina:", "").strip()

        # Extração da tabela de alunos
        linhas = soup.find_all("tr")
        i = 0
        while i < len(linhas):
            cols = linhas[i].find_all(["td", "th"])
            textos = [c.text.strip() for c in cols]
            
            if textos and textos[0].isdigit():
                num = int(textos[0])
                ra = textos[1] if len(textos) > 1 else ""
                div_nome = linhas[i].find("div", class_="truncate")
                nome = div_nome.text.strip() if div_nome else (textos[2].split("Ingresso")[0].strip() if len(textos) > 2 else "")
                status = textos[-1] if len(textos) >= 4 else ""

                if ra not in dados_alunos:
                    dados_alunos[ra] = {"Nº": num, "RA": ra, "NOME": nome, "BIMESTRES": {}}

                dados_alunos[ra]["BIMESTRES"][bimestre] = status
            i += 1

    return dados_alunos, info_cabecalho, bimestres_encontrados


def gerar_excel_fluencia(dados_alunos, info_cabecalho, caminho_saida):
    """Gera o arquivo Excel formatado com os dados consolidados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acompanhamento Anual"
    ws.views.sheetView[0].showGridLines = True

    # Configurações de página para impressão A4 Retrato
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left, ws.page_margins.right = 0.4, 0.4
    ws.page_margins.top, ws.page_margins.bottom = 0.5, 0.5

    # Cabeçalho Principal
    ws.merge_cells("A1:G1")
    ws["A1"] = "SECRETARIA MUNICIPAL DE EDUCAÇÃO DE ILHABELA"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:G2")
    ws["A2"] = f"RELATÓRIO DE ACOMPANHAMENTO DA FLUÊNCIA LEITORA — {info_cabecalho.get('Ano Letivo', '2026')}"
    ws["A2"].font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    fill_quadro = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    ws["A3"], ws["B3"], ws["D3"], ws["E3"] = "Escola:", info_cabecalho.get("Escola", ""), "Turma:", info_cabecalho.get("Turma", "")
    ws["A4"], ws["B4"], ws["D4"], ws["E4"] = "Série:", info_cabecalho.get("Série", ""), "Disciplina:", info_cabecalho.get("Disciplina", "")

    for r in [3, 4]:
        ws.row_dimensions[r].height = 18
        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            ws[f"{col}{r}"].fill = fill_quadro
            ws[f"{col}{r}"].font = Font(name="Calibri", size=9, bold=(col in ["A", "D"]), color="1E293B")

    headers = ["Nº", "RA", "Nome do Aluno", "1º Bim", "2º Bim", "3º Bim", "4º Bim"]
    fill_header = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    for idx, h in enumerate(headers, 1):
        c = ws.cell(row=6, column=idx, value=h)
        c.fill = fill_header
        c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center" if idx != 3 else "left", vertical="center")

    thin_border = Border(left=Side(style="thin", color="E2E8F0"), right=Side(style="thin", color="E2E8F0"), top=Side(style="thin", color="E2E8F0"), bottom=Side(style="thin", color="E2E8F0"))

    row_idx = 7
    for ra, al in sorted(dados_alunos.items(), key=lambda x: x[1]["Nº"]):
        ws.row_dimensions[row_idx].height = 20
        c_num = ws.cell(row=row_idx, column=1, value=al["Nº"])
        c_ra = ws.cell(row=row_idx, column=2, value=al["RA"])
        c_nome = ws.cell(row=row_idx, column=3, value=al["NOME"])

        for c in [c_num, c_ra, c_nome]:
            c.font = Font(name="Calibri", size=9.5, color="1F2937")
            c.border = thin_border
        c_num.alignment = Alignment(horizontal="center", vertical="center")
        c_ra.alignment = Alignment(horizontal="center", vertical="center")
        c_nome.alignment = Alignment(horizontal="left", vertical="center")

        for col_b, bim_name in enumerate(["1º Bimestre", "2º Bimestre", "3º Bimestre", "4º Bimestre"], start=4):
            sigla = al["BIMESTRES"].get(bim_name, "-")
            c_bim = ws.cell(row=row_idx, column=col_b, value=sigla)
            c_bim.border = thin_border
            c_bim.alignment = Alignment(horizontal="center", vertical="center")

            if sigla in MAPA_DICT:
                _, _, bg, fg = MAPA_DICT[sigla]
                c_bim.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                c_bim.font = Font(name="Calibri", size=9, bold=True, color=fg)
            else:
                c_bim.font = Font(name="Calibri", size=9, color="94A3B8")
        row_idx += 1

    # Legenda Completa
    row_idx += 1
    ws.cell(row=row_idx, column=1, value="LEGENDA DE NÍVEIS DE FLUÊNCIA LEITORA:").font = Font(name="Calibri", size=9.5, bold=True, color="1E293B")

    for sigla, desc, bg, fg in LEGENDAS_COMPLETAS:
        row_idx += 1
        ws.row_dimensions[row_idx].height = 18
        c_badge = ws.cell(row=row_idx, column=1, value=sigla)
        c_badge.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        c_badge.font = Font(name="Calibri", size=8.5, bold=True, color=fg)
        c_badge.alignment = Alignment(horizontal="center", vertical="center")
        c_badge.border = thin_border

        ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=4)
        c_desc = ws.cell(row=row_idx, column=2, value=f"- {desc}")
        c_desc.font = Font(name="Calibri", size=8.5, bold=True, color="334155")
        c_desc.alignment = Alignment(horizontal="left", vertical="center")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10

    wb.save(caminho_saida)